"""drainage-report（排水計畫書本文）：範例 build 全綠 + report_check 故意錯（漏章、圖名錯、數值錯、
刪圖、md / docx 不一致、合規宣稱、移除 verified: false 提醒…）。不呼叫 LLM。"""

from __future__ import annotations

import copy
import json
import re
import shutil
import sys
from pathlib import Path

import pytest
from pydantic import ValidationError

pytest.importorskip("docx")

from draftagent.build import build, load_spec
from draftagent.report.model import DrainageReportSpec
from draftagent.verify.report_check import check_report

ROOT = Path(__file__).resolve().parents[1]
EXAMPLE = ROOT / "examples" / "drainage-report.json"


@pytest.fixture(scope="module")
def built(tmp_path_factory):
    res = build(load_spec(EXAMPLE), out=tmp_path_factory.mktemp("rep"), preview=False)
    return res


def test_example_build_green(built):
    assert built.ok, built.report()
    names = [r.name for r in built.results]
    assert names[-1] == "report_check" and "drainage_plan_check" in names
    rep = built.results[-1]
    assert rep.metrics["figures"] == {"md": 5, "docx": 5, "expected": 5}
    assert rep.metrics["values_compared"] > 200


def test_docx_readback(built):
    import docx

    d = docx.Document(str(built.artifacts["report_docx"]))
    heads = [p.text for p in d.paragraphs if p.style.name == "Heading 1"]
    assert heads[0] == "摘要" and heads[-1] == "附圖" and len(heads) == 8
    assert len(d.tables) == 12
    assert len(d.inline_shapes) == 5
    s = d.sections[0]
    assert round(s.page_width.mm) == 210 and round(s.page_height.mm) == 297


# ---------------------------------------------------------------- 竄改工具


def _copy(built, tmp_path) -> tuple[DrainageReportSpec, dict[str, Path]]:
    src = built.run_dir
    dst = tmp_path / "run"
    shutil.copytree(src, dst)
    arts = {k: dst / Path(v).relative_to(src) for k, v in built.artifacts.items()}
    return load_spec(arts["spec"]), arts


def _edit(arts, old: str, new: str, md: bool = True, docx_: bool = True, count: int = 1) -> None:
    import docx

    if md:
        p = arts["report_md"]
        s = p.read_text(encoding="utf-8")
        assert old in s, old
        p.write_text(s.replace(old, new, count if count > 0 else -1), encoding="utf-8")
    if docx_:
        d = docx.Document(str(arts["report_docx"]))
        n = 0
        paras = list(d.paragraphs) + [
            p for t in d.tables for r in t.rows for c in r.cells for p in c.paragraphs
        ]
        for para in paras:
            for run in para.runs:
                if old in run.text and (count <= 0 or n < count):
                    run.text = run.text.replace(old, new)
                    n += 1
        assert n, f"docx 沒有「{old}」"
        d.save(str(arts["report_docx"]))


def _errors(spec, arts) -> str:
    r = check_report(spec, arts)
    assert not r.ok
    return "\n".join(r.errors)


def test_missing_chapter(built, tmp_path):
    import docx

    spec, arts = _copy(built, tmp_path)
    p = arts["report_md"]
    p.write_text(p.read_text(encoding="utf-8").replace("## 一、目的\n\n", ""), encoding="utf-8")
    d = docx.Document(str(arts["report_docx"]))
    h = next(x for x in d.paragraphs if x.text == "一、目的")
    h._p.getparent().remove(h._p)
    d.save(str(arts["report_docx"]))
    e = _errors(spec, arts)
    assert "章節不符" in e and "一、目的" in e
    assert "不一致" not in e  # 兩邊同改：只報章節


def test_wrong_sheet_name(built, tmp_path):
    spec, arts = _copy(built, tmp_path)
    _edit(arts, "附圖 1 排水系統平面設計圖", "附圖 1 排水平面圖")
    _edit(arts, "| 6 | 排水系統平面設計圖 |", "| 6 | 排水平面圖 |", docx_=False)
    e = _errors(spec, arts)
    assert "附圖 1 圖說" in e and "統一圖名" in e


def test_wrong_drawing_no(built, tmp_path):
    spec, arts = _copy(built, tmp_path)
    _edit(
        arts, "附圖 5 排水系統集水區分析圖（圖號 DR-02", "附圖 5 排水系統集水區分析圖（圖號 DR-09"
    )
    assert "附圖 5 圖說" in _errors(spec, arts)


def test_wrong_expression_value(built, tmp_path):
    spec, arts = _copy(built, tmp_path)
    _edit(arts, "提供貯集量 V_prov = 124.80", "提供貯集量 V_prov = 125.80")
    assert "V_prov" in _errors(spec, arts)


def test_wrong_table_value(built, tmp_path):
    spec, arts = _copy(built, tmp_path)
    # 表 5-2 P1 的流速
    _edit(arts, "1.140", "1.240")
    assert "表 5-2 管渠水理計算 P1 流速" in _errors(spec, arts)


def test_rounding_must_be_correct(built, tmp_path):
    spec, arts = _copy(built, tmp_path)
    _edit(arts, "0.029018", "0.029019")  # 摘要 Q_out 差 1e-6 也要抓
    assert "Q_out" in _errors(spec, arts)


def test_unknown_expression_label(built, tmp_path):
    spec, arts = _copy(built, tmp_path)
    _edit(arts, "適用法規路徑由工程師判定。", "適用法規路徑由工程師判定。安全係數 = 1.2。")
    assert "無法追溯" in _errors(spec, arts)


def test_deleted_figure(built, tmp_path):
    import docx

    spec, arts = _copy(built, tmp_path)
    p = arts["report_md"]
    s = p.read_text(encoding="utf-8")
    s = re.sub(r"!\[附圖 3\]\([^)]+\)\n\n附圖 3 [^\n]+\n\n", "", s)
    p.write_text(s, encoding="utf-8")
    d = docx.Document(str(arts["report_docx"]))
    pics = [
        x
        for x in d.paragraphs
        if x._p.findall(".//{http://schemas.openxmlformats.org/drawingml/2006/main}blip")
    ]
    cap = next(x for x in d.paragraphs if x.text.startswith("附圖 3 "))
    for x in (pics[2], cap):
        x._p.getparent().remove(x._p)
    d.save(str(arts["report_docx"]))
    assert "附圖數不符" in _errors(spec, arts)


def test_blank_figure(built, tmp_path):
    from PIL import Image

    spec, arts = _copy(built, tmp_path)
    fig = arts["report_md"].parent / "report_figures" / "fig-04.png"
    Image.new("RGB", (800, 600), "white").save(fig)
    assert "空白" in _errors(spec, arts)


def test_md_docx_mismatch(built, tmp_path):
    spec, arts = _copy(built, tmp_path)
    _edit(arts, "基地範圍、建物外框", "基地範圍與建物外框", docx_=False)
    assert "不一致" in _errors(spec, arts)


def test_compliance_claim(built, tmp_path):
    spec, arts = _copy(built, tmp_path)
    _edit(arts, "適用法規路徑由工程師判定。", "適用法規路徑由工程師判定。本設計符合法規。")
    assert "合規宣稱" in _errors(spec, arts)


def test_negated_compliance_is_fine(built):
    from draftagent.verify.report_check import compliance_claims

    assert compliance_claims("通過檢查項目，並不代表符合法規。") == []
    assert compliance_claims("不能寫「本設計符合法規」") == []
    assert compliance_claims("本設計符合法規。")


def test_removed_verified_false(built, tmp_path):
    spec, arts = _copy(built, tmp_path)
    _edit(arts, "verified: false", "", count=0)
    e = _errors(spec, arts)
    assert "verified: false 提醒" in e and "未註明" in e and "核對狀態" in e


def test_missing_conclusion(built, tmp_path):
    spec, arts = _copy(built, tmp_path)
    _edit(arts, "並不代表符合法規。", "。")
    assert "固定結論句" in _errors(spec, arts)


def test_flood_potential_only_risk_hint(built, tmp_path):
    spec, arts = _copy(built, tmp_path)
    _edit(
        arts,
        "適用法規路徑由工程師判定。",
        "適用法規路徑由工程師判定。依淹水潛勢圖基地不淹水，可減設滯洪。",
    )
    assert "淹水潛勢" in _errors(spec, arts)


def test_client_id_pattern(built, tmp_path):
    spec, arts = _copy(built, tmp_path)
    firm = "某某工程" + "顧問有限公司"  # 拆開：避免本檔被 check_no_client_ids 命中
    _edit(arts, "設計單位：○○○", f"設計單位：{firm}")
    assert "識別資訊" in _errors(spec, arts)


def test_regulatory_value_needs_yaml_tag(built, tmp_path):
    spec, arts = _copy(built, tmp_path)
    p = arts["report_md"]
    s = p.read_text(encoding="utf-8")
    old = re.search(r"最小貯集量係數 k_TY = 0\.051 m³/m²〔[^〕]+〕", s).group(0)
    _edit(arts, old, "最小貯集量係數 k_TY = 0.051 m³/m²")
    assert "旁缺 yaml 來源標註" in _errors(spec, arts)


def test_patterns_in_sync():
    import importlib.util

    from draftagent.eval.tasks import COMPLIANCE_CLAIM as EVAL_CLAIM
    from draftagent.verify import report_check as rc

    sp = importlib.util.spec_from_file_location("cnc", ROOT / "scripts" / "check_no_client_ids.py")
    mod = importlib.util.module_from_spec(sp)
    sys.modules["cnc"] = mod
    try:
        sp.loader.exec_module(mod)
    finally:
        sys.modules.pop("cnc", None)
    assert [(k, p.pattern) for k, p in mod.PATTERNS] == [
        (k, p.pattern) for k, p in rc.CLIENT_ID_PATTERNS
    ]
    assert rc.COMPLIANCE_CLAIM == EVAL_CLAIM


# ---------------------------------------------------------------- spec / 照片 / 缺單價 / 字型


def _raw() -> dict:
    return json.loads(EXAMPLE.read_text(encoding="utf-8"))


def test_notes_must_not_contain_digits():
    d = _raw()
    d["report"]["notes"] = {"summary": ["提供貯集量為 130 m³，大於必要量。"]}
    with pytest.raises(ValidationError, match="不得含數字"):
        DrainageReportSpec.model_validate(d)


def test_unit_price_must_match_quantity_row():
    d = _raw()
    d["report"]["unit_prices"].append({"item": "排水管", "unit_price": 1})  # 兩種規格，未指定 spec
    with pytest.raises(ValidationError, match="對應數量表"):
        DrainageReportSpec.model_validate(d)


def test_report_with_photo_and_missing_price(tmp_path, monkeypatch):
    from PIL import Image, ImageDraw

    photo = tmp_path / "site.png"
    im = Image.new("RGB", (640, 480), "white")
    ImageDraw.Draw(im).rectangle((100, 100, 500, 380), fill="gray")
    im.save(photo)
    d = _raw()
    d["report"]["photos"] = [{"path": str(photo), "caption": "基地南側道路與既有側溝（合成照片）"}]
    d["report"]["unit_prices"] = [u for u in d["report"]["unit_prices"] if u["item"] != "人孔"]
    res = build(DrainageReportSpec.model_validate(d), out=tmp_path / "b", preview=False)
    assert res.ok, res.report()
    rep = res.results[-1]
    assert rep.metrics["figures"]["expected"] == 6
    assert any("人孔" in w and "未提供單價" in w for w in rep.warnings)
    md = Path(res.artifacts["report_md"]).read_text(encoding="utf-8")
    assert "總工程造價（不含未提供單價項目） = " in md
    assert "附圖 1 基地現況照片：" in md and "| 附圖 1 | site.png |" in md


def test_font_fallback():
    from draftagent.report.vocab import load_format
    from draftagent.report.writers import pick_east_asia_font

    lay = copy.deepcopy(load_format().layout)
    assert pick_east_asia_font(lay, {"DFKai-SB"}) == ("標楷體", [])
    name, w = pick_east_asia_font(lay, {"Kaiti TC"})
    assert name == "Kaiti TC" and w
    name, w = pick_east_asia_font(lay, set())
    assert name == "標楷體" and "未安裝" in w[0]
