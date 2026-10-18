"""內容區塊 → report.md / report.docx（python-docx）。兩者由同一份區塊序列產生，段落文字相同。

docx 版面（A4、邊界、字型、字號、行距）讀 `report/codes/report_format.yaml` `layout`（verified: false，
版面數值為 repo 暫定）。東亞字型本機找不到時依 yaml fallback 並回傳 WARNING。純函式，不得呼叫 LLM。
"""

from __future__ import annotations

import os
from pathlib import Path
from typing import Any

from draftagent.report.content import Block

# ---------------------------------------------------------------- markdown


def write_md(blocks: list[Block], path: Path) -> Path:
    path = Path(path)
    out: list[str] = []

    def para(s: str) -> None:
        if out and out[-1] != "":
            out.append("")
        out.append(s)

    prev = None
    for b in blocks:
        if b.type == "title":
            para(f"# {b.text}")
        elif b.type == "h1":
            para(f"## {b.text}")
        elif b.type == "h2":
            para(f"### {b.text}")
        elif b.type == "li":
            if prev == "li":
                out.append(f"- {b.text}")
            else:
                para(f"- {b.text}")
        elif b.type == "table":
            para("| " + " | ".join(b.header) + " |")
            out.append("|" + "|".join("---" for _ in b.header) + "|")
            out += ["| " + " | ".join(r) + " |" for r in b.rows]
        elif b.type == "image":
            rel = os.path.relpath(b.path, path.parent)
            para(f"![{b.text}]({Path(rel).as_posix()})")
        else:  # p / caption / toc
            para(b.text)
        prev = b.type
    path.write_text("\n".join(out) + "\n", encoding="utf-8")
    return path


# ---------------------------------------------------------------- docx


def installed_font_families() -> set[str]:
    from matplotlib import font_manager

    names = {f.name for f in font_manager.fontManager.ttflist}
    for f in font_manager.fontManager.ttflist:
        names.add(Path(f.fname).stem)
    return names


def pick_east_asia_font(
    layout: dict[str, Any], installed: set[str] | None = None
) -> tuple[str, list[str]]:
    """(字型名, warnings)。標楷體（或其別名）已安裝 → 用標楷體；否則依 fallbacks。"""
    fonts = layout["fonts"]
    want = fonts["east_asia"]
    installed = installed_font_families() if installed is None else installed
    if any(a in installed for a in fonts.get("east_asia_aliases", [want])):
        return want, []
    for fb in fonts.get("east_asia_fallbacks", []):
        if fb in installed:
            return fb, [
                f"字型「{want}」本機未安裝，docx 東亞字型改用「{fb}」（送審前請改回{want}）"
            ]
    return want, [
        f"字型「{want}」與 fallback {fonts.get('east_asia_fallbacks', [])} 本機皆未安裝；"
        f"docx 仍指定「{want}」，開啟端須有此字型"
    ]


def _set_fonts(
    style_or_run, latin: str, east: str, size_pt: float | None, bold: bool | None = None
) -> None:
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    font = style_or_run.font
    font.name = latin
    if size_pt is not None:
        font.size = Pt(size_pt)
    if bold is not None:
        font.bold = bold
    font.color.rgb = RGBColor(0, 0, 0)
    rpr = style_or_run.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.insert(0, rfonts)
    for att in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if rfonts.get(qn(att)) is not None:
            del rfonts.attrib[qn(att)]
    rfonts.set(qn("w:ascii"), latin)
    rfonts.set(qn("w:hAnsi"), latin)
    rfonts.set(qn("w:cs"), latin)
    rfonts.set(qn("w:eastAsia"), east)


def _toc(paragraph, placeholder: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    def fld(kind: str):
        r = OxmlElement("w:r")
        f = OxmlElement("w:fldChar")
        f.set(qn("w:fldCharType"), kind)
        r.append(f)
        return r

    p = paragraph._p
    p.append(fld("begin"))
    r = OxmlElement("w:r")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = 'TOC \\o "1-2" \\h \\z \\u'
    r.append(instr)
    p.append(r)
    p.append(fld("separate"))
    paragraph.add_run(placeholder)
    p.append(fld("end"))


def write_docx(
    blocks: list[Block], path: Path, layout: dict[str, Any]
) -> tuple[Path, dict[str, Any]]:
    """回傳 (路徑, meta)；meta = {east_asia_font, latin_font, warnings}。"""
    try:
        import docx
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Mm, Pt
    except ImportError as e:  # pragma: no cover - 相依未安裝
        raise RuntimeError("產生 report.docx 需要 python-docx（uv pip install python-docx）") from e

    east, warnings = pick_east_asia_font(layout)
    latin = layout["fonts"]["latin"]
    sz = layout["sizes_pt"]
    page, mg = layout["page"], layout["margins_mm"]

    d = docx.Document()
    sec = d.sections[0]
    sec.page_width, sec.page_height = Mm(page["width_mm"]), Mm(page["height_mm"])
    sec.top_margin, sec.bottom_margin = Mm(mg["top"]), Mm(mg["bottom"])
    sec.left_margin, sec.right_margin = Mm(mg["left"]), Mm(mg["right"])

    styles = d.styles
    for name, size, bold in (
        ("Normal", sz["body"], None),
        ("Title", sz["title"], True),
        ("Heading 1", sz["chapter"], True),
        ("Heading 2", sz["section"], True),
        ("Caption", sz["caption"], False),
        ("List Bullet", sz["body"], None),
    ):
        st = styles[name]
        _set_fonts(st, latin, east, size, bold)
        if name in ("Normal", "List Bullet"):
            st.paragraph_format.line_spacing = layout["line_spacing"]
        if name == "Caption":
            st.font.italic = False

    fig_w = Mm(layout["figure_width_mm"])
    for b in blocks:
        if b.type == "title":
            p = d.add_paragraph(b.text, style="Title")
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif b.type in ("h1", "h2"):
            p = d.add_heading(b.text, level=1 if b.type == "h1" else 2)
        elif b.type == "li":
            p = d.add_paragraph(b.text, style="List Bullet")
        elif b.type == "caption":
            p = d.add_paragraph(b.text, style="Caption")
        elif b.type == "toc":
            p = d.add_paragraph()
            _toc(p, b.text)
        elif b.type == "table":
            t = d.add_table(rows=1 + len(b.rows), cols=len(b.header))
            t.style = "Table Grid"
            for i, row in enumerate([b.header, *b.rows]):
                for j, txt in enumerate(row):
                    cell = t.cell(i, j)
                    cell.text = txt
                    for run in cell.paragraphs[0].runs:
                        run.font.size = Pt(sz["table"])
                        run.font.bold = i == 0
            continue
        elif b.type == "image":
            p = d.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            from PIL import Image

            with Image.open(b.path) as im:
                w_px, h_px = im.size
            max_h = Mm(page["height_mm"] - mg["top"] - mg["bottom"] - 30)  # 留圖說空間
            w = min(fig_w, int(max_h * w_px / h_px)) if h_px else fig_w
            p.add_run().add_picture(str(b.path), width=w)
        else:
            p = d.add_paragraph(b.text)
        if b.page_break:
            p.paragraph_format.page_break_before = True
    d.core_properties.title = next((b.text for b in blocks if b.type == "title"), "")
    d.save(str(path))
    return Path(path), {"east_asia_font": east, "latin_font": latin, "warnings": warnings}
